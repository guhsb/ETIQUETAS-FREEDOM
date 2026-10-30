# -*- coding: utf-8 -*-
"""Histórico de etiquetas impressas: base em CSV (segura) + Excel gerado a partir dela."""
import csv
import datetime as dt
import threading
import getpass
import os
import re
import socket
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from etiqueta_core import pasta_dados

COLS = ["DATA", "HORA", "NF", "VOLUME", "TOTAL", "COD CLIENTE", "CLIENTE", "CIDADE", "UF",
        "TRANSPORTADORA", "REMETENTE", "IMPRESSORA", "USUARIO", "STATUS", "JOB"]
COLS_VISIVEIS = [c for c in COLS if c != "JOB"]          # JOB é só para o programa atualizar a situação
TITULOS = {"DATA": "Data", "HORA": "Hora", "NF": "NF", "VOLUME": "Volume", "TOTAL": "De",
           "COD CLIENTE": "Cód. cliente", "CLIENTE": "Cliente", "CIDADE": "Cidade", "UF": "UF",
           "TRANSPORTADORA": "Transportadora", "REMETENTE": "Remetente", "IMPRESSORA": "Impressora",
           "USUARIO": "Usuário", "STATUS": "Situação", "JOB": "Trabalho"}
LARGURAS = {"DATA": 12, "HORA": 9, "NF": 11, "VOLUME": 9, "TOTAL": 7, "COD CLIENTE": 13, "CLIENTE": 46,
            "CIDADE": 24, "UF": 6, "TRANSPORTADORA": 24, "REMETENTE": 26, "IMPRESSORA": 28, "USUARIO": 22, "STATUS": 16, "JOB": 10}
NOME_EXCEL = "HISTORICO_ETIQUETAS.xlsx"
_TRAVA = threading.RLock()      # gravar e atualizar a base nunca ao mesmo tempo (tela e segundo plano)


def caminho_csv():
    return os.path.join(pasta_dados(), "historico_etiquetas.csv")


def _sem_acento(t):
    return "".join(c for c in unicodedata.normalize("NFD", str(t)) if unicodedata.category(c) != "Mn").upper()


def nova_linha(nf, volume, total, cli, transportadora, remetente, impressora, agora=None, job="", status="ENVIADA"):
    agora = agora or dt.datetime.now()
    try:
        usuario = f"{getpass.getuser()}@{socket.gethostname()}"
    except Exception:
        usuario = ""
    return {"DATA": agora.date().isoformat(), "HORA": agora.strftime("%H:%M:%S"), "NF": str(nf),
            "VOLUME": str(volume), "TOTAL": str(total), "COD CLIENTE": cli["cod"], "CLIENTE": cli["nome"],
            "CIDADE": cli["cidade"], "UF": cli["uf"], "TRANSPORTADORA": transportadora,
            "REMETENTE": remetente, "IMPRESSORA": impressora, "USUARIO": usuario, "STATUS": status, "JOB": job}


def _ajusta_cabecalho():
    """Arquivo criado por uma versão antiga (sem as colunas novas): reescreve com o cabeçalho atual."""
    arq = caminho_csv()
    if not os.path.exists(arq):
        return
    with open(arq, encoding="utf-8-sig") as f:
        primeira = f.readline().strip().split(";")
    if primeira != COLS:
        df = ler()
        df["STATUS"] = df["STATUS"].replace("", "IMPRESSA")           # o que já existia foi considerado impresso
        _grava(df)


def _grava(df):
    arq = caminho_csv()
    tmp = arq + ".tmp"
    df[COLS].to_csv(tmp, sep=";", index=False, encoding="utf-8-sig")
    os.replace(tmp, arq)


def registrar(linhas):
    """Acrescenta as linhas na base (CSV). Nunca fica 'travada' como um Excel aberto."""
    with _TRAVA:
        _ajusta_cabecalho()
        arq = caminho_csv()
        novo = not os.path.exists(arq)
        with open(arq, "a", newline="", encoding="utf-8-sig" if novo else "utf-8") as f:
            w = csv.DictWriter(f, COLS, delimiter=";")
            if novo:
                w.writeheader()
            for l in linhas:
                w.writerow({c: l.get(c, "") for c in COLS})


def marca_status(job, volumes, status):
    """Atualiza a situação das etiquetas de um trabalho (ex.: ENVIADA -> IMPRESSA ou NÃO IMPRESSA)."""
    if not job:
        return 0
    with _TRAVA:
        df = ler()
        sel = (df["JOB"] == job) & df["VOLUME"].isin([str(v) for v in volumes])
        if sel.any():
            df.loc[sel, "STATUS"] = status
            _grava(df)
        return int(sel.sum())


def ler():
    try:
        df = pd.read_csv(caminho_csv(), sep=";", dtype=str, encoding="utf-8-sig").fillna("")
    except (FileNotFoundError, pd.errors.EmptyDataError):
        df = pd.DataFrame(columns=COLS)
    for c in COLS:
        if c not in df.columns:
            df[c] = ""
    return df[COLS]


def data_br(texto):
    """'20/10/2026', '20/10/26', '20/10' ou '20102026' -> date. Vazio -> None."""
    t = (texto or "").strip()
    if not t:
        return None
    t = re.sub(r"[.\-]", "/", t)
    if re.fullmatch(r"\d{8}", t):
        t = f"{t[:2]}/{t[2:4]}/{t[4:]}"
    p = t.split("/")
    try:
        d, m = int(p[0]), int(p[1])
        a = int(p[2]) if len(p) > 2 and p[2] else dt.date.today().year
        if a < 100:
            a += 2000
        return dt.date(a, m, d)
    except (ValueError, IndexError):
        raise ValueError(f"Data inválida: “{texto}”. Use dia/mês/ano, por exemplo 20/10/2026.")


def filtrar(df, de=None, ate=None, transportadora="", nf="", cliente="", texto=""):
    r = df
    if de:
        r = r[r["DATA"] >= de.isoformat()]
    if ate:
        r = r[r["DATA"] <= ate.isoformat()]
    if transportadora and transportadora.upper() not in ("TODAS", "(TODAS)"):
        r = r[r["TRANSPORTADORA"].str.upper() == transportadora.upper()]
    if nf.strip():
        r = r[r["NF"].str.contains(re.escape(nf.strip()), regex=True)]
    if cliente.strip():
        q = _sem_acento(cliente.strip())
        alvo = (r["COD CLIENTE"] + " " + r["CLIENTE"]).map(_sem_acento)
        r = r[alvo.str.contains(re.escape(q), regex=True)]
    if texto.strip():
        q = _sem_acento(texto.strip())
        r = r[r.apply(lambda x: q in _sem_acento(" ".join(x.astype(str))), axis=1)]
    return r.sort_values(["DATA", "HORA"], ascending=False, kind="stable")


def resumo(df):
    return {"etiquetas": len(df), "nfs": df["NF"].nunique(), "clientes": df["COD CLIENTE"].nunique()}


def data_exibicao(iso):
    try:
        return dt.date.fromisoformat(iso).strftime("%d/%m/%Y")
    except ValueError:
        return iso


def exportar_xlsx(df, caminho):
    """Gera o Excel formatado (padrão Freedom), em modo de gravação direta (rápido mesmo com milhares de linhas).
    Levanta PermissionError se o arquivo estiver aberto."""
    from openpyxl.cell import WriteOnlyCell
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Etiquetas")
    azul, zebra, linha_cor = "0F1B3D", "F6F8FC", "E1E7EF"
    for i, c in enumerate(COLS_VISIVEIS, 1):
        ws.column_dimensions[get_column_letter(i)].width = LARGURAS[c]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS_VISIVEIS))}{max(2, len(df) + 1)}"
    f_cab = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    p_cab = PatternFill("solid", fgColor=azul)
    centro = Alignment(horizontal="center", vertical="center")
    cab = []
    for c in COLS_VISIVEIS:
        cel = WriteOnlyCell(ws, value=TITULOS[c])
        cel.font, cel.fill, cel.alignment = f_cab, p_cab, centro
        cab.append(cel)
    ws.append(cab)
    f_corpo = Font(name="Arial", size=10, color="2B3550")          # estilos criados uma vez só
    p_zebra = PatternFill("solid", fgColor=zebra)
    borda = Border(bottom=Side(style="thin", color=linha_cor))
    al_centro = Alignment(horizontal="center")
    centradas = {i for i, c in enumerate(COLS_VISIVEIS) if c in ("DATA", "HORA", "NF", "VOLUME", "TOTAL", "COD CLIENTE", "UF", "STATUS")}
    i_data = COLS_VISIVEIS.index("DATA")
    numericas = {i for i, c in enumerate(COLS_VISIVEIS) if c in ("NF", "VOLUME", "TOTAL", "COD CLIENTE")}
    for n, linha in enumerate(df[COLS_VISIVEIS].values.tolist(), start=2):
        try:
            linha[i_data] = dt.date.fromisoformat(linha[i_data])
        except (ValueError, TypeError):
            pass
        listrada = n % 2 == 1
        cels = []
        for i, v in enumerate(linha):
            if i in numericas and str(v).isdigit():
                v = int(v)
            cel = WriteOnlyCell(ws, value=v)
            cel.font, cel.border = f_corpo, borda
            if listrada:
                cel.fill = p_zebra
            if i == i_data:
                cel.number_format = "dd/mm/yyyy"
            if i in centradas:
                cel.alignment = al_centro
            cels.append(cel)
        ws.append(cels)
    tmp = caminho + ".tmp"
    wb.save(tmp)
    os.replace(tmp, caminho)


def atualizar_excel(caminho):
    """Regera o Excel completo a partir da base. Devolve (ok, mensagem_de_erro)."""
    try:
        exportar_xlsx(ler().sort_values(["DATA", "HORA"], kind="stable"), caminho)
        return True, ""
    except PermissionError:
        return False, "O arquivo está aberto no Excel. Feche-o para atualizar (os dados estão salvos e entram na próxima impressão)."
    except Exception as ex:
        return False, str(ex)

